"""Generate the illustrated portfolio from the same catalog as the website."""
import hashlib
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / 'assets/project_descriptions/fullportfolio.pptx'
STAMP = ROOT / 'scripts/presentation-build.json'
catalog = json.loads((ROOT / 'assets/data/projects.json').read_text(encoding='utf-8'))
inputs = [ROOT / 'assets/data/projects.json', Path(__file__).resolve()]
inputs += [ROOT / m['src'] for p in catalog for m in p['media'] if m['type'] == 'image']
inputs += [(ROOT / p['supporting_page']).parent / 'project-notes.md' for p in catalog]
fingerprints = {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs}
if '--check' in sys.argv:
    saved = json.loads(STAMP.read_text())
    assert saved['inputs'] == fingerprints, 'Presentation inputs changed; rebuild projects.'
    assert saved['output'] == hashlib.sha256(TARGET.read_bytes()).hexdigest(), 'Presentation changed; rebuild.'
    print('PASS: presentation matches current project content and images.')
    raise SystemExit(0)

sys.path.insert(0, str(ROOT / '.tmp-presentation-tools'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from PIL import Image, ImageOps
from io import BytesIO

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG, WHITE, MUTED, ACCENT = 'FFFFFF', '1F1F1F', '666666', 'ED7D31'


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False, compact=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = 'Aptos'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(2 if compact else 6)
    return shape


def photo(slide, item, x, y, w, h):
    with Image.open(ROOT / item['src']) as source:
        im = ImageOps.exif_transpose(source).convert('RGB')
        im.thumbnail((1800, 1400))
        ratio = min(w / im.width, h / im.height)
        width, height = im.width * ratio, im.height * ratio
        buffer = BytesIO()
        im.save(buffer, format='JPEG', quality=90)
    slide.shapes.add_picture(buffer, Inches(x + (w-width)/2), Inches(y + (h-height)/2), width=Inches(width), height=Inches(height))
    text(slide, item['caption'], x, y+h+.08, w, .48, 11, MUTED)


def base(project):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    text(slide, f"{project['year']}  /  {project['context']}", .55, .3, 12.2, .4, 13, ACCENT)
    text(slide, project['title'], .55, .85, 12.2, .68, 29, bold=True)
    footer = text(slide, 'Henry Luo  |  Engineering Portfolio  |  View project and supporting files', .55, 7.05, 11.9, .25, 10, MUTED)
    footer.text_frame.paragraphs[0].runs[0].hyperlink.address = 'https://the-other-hello-there.github.io/HenryLuo-Portfolio/#project-' + project['id']
    text(slide, str(len(prs.slides)), 12.2, 7.05, .5, .25, 10, MUTED)
    notes = '\n\n'.join(f'{key}: {value}' for key, value in project.items() if key not in {'media', 'supporting_page'})
    notes += '\n\n' + ((ROOT / project['supporting_page']).parent / 'project-notes.md').read_text(encoding='utf-8')
    notes += '\n\nImages and demonstrations:\n' + '\n'.join(m['caption'] + ': ' + m['src'] for m in project['media'])
    slide.notes_slide.notes_text_frame.text = notes
    return slide


TEAM_PROJECTS = {
    'feline': 3,
    'battle-bot': 3,
    'ghost-arm': 2,
    'infill-thermal-conduction': 4,
    'headwind-safety': 3,
    'hrc-l2-leg': 4,
    'rc-car': 3,
}


def section_content(project):
    bullets = project['presentation_bullets']
    if project['id'] in TEAM_PROJECTS:
        contribution_count = TEAM_PROJECTS[project['id']]
        return [
            ('Goals', [project['objective'], project['challenges'][0]]),
            ('Contributions', bullets[:contribution_count]),
            ('Achievements', bullets[contribution_count:]),
        ]
    if 'Personal Project' in project['tags']:
        return [('Achievements', bullets)]
    return [
        ('Goals', [project['objective'], project['challenges'][0]]),
        ('Achievements', bullets),
    ]


def add_sections(slide, project):
    sections = section_content(project)
    section_height = 4.55 / len(sections)
    for index, (heading, bullets) in enumerate(sections):
        y = 2.05 + index * section_height
        text(slide, heading, .65, y, 6.1, .25, 13, ACCENT, bold=True)
        text(slide, '\n'.join('\u2022 ' + line for line in bullets), .65, y + .3, 6.1, section_height - .32, 11.5, WHITE, compact=True)


for project in catalog:
    images = [m for m in project['media'] if m['type'] == 'image']
    slide = base(project)
    add_sections(slide, project)
    chosen = images[:2]
    if project['id'] == 'infill-thermal-conduction': chosen = [images[0], images[-1]]
    if project['id'] == 'sunglasses-holder': chosen = [images[0], images[-2]]
    if project['id'] == 'feline': chosen = [images[0], images[3]]
    if project['id'] == 'clothes-hanger': chosen = [images[1]]
    if len(chosen) == 1:
        photo(slide, chosen[0], 7.2, 2.2, 5.55, 4.05)
    else:
        for i, item in enumerate(chosen):
            photo(slide, item, 7.2, 2.1+i*2.35, 5.55, 1.8)
prs.core_properties.title = 'Henry Luo — Engineering Project Portfolio'
prs.core_properties.author = 'Henry Luo'
prs.core_properties.subject = 'Project objectives, contributions, evidence, and limitations'
# Keep hyperlink text aligned with the orange presentation theme.
from lxml import etree
for part in prs.part.package.iter_parts():
    if str(part.partname).startswith('/ppt/theme/'):
        theme = etree.fromstring(part.blob)
        for name in ('hlink', 'folHlink'):
            for node in theme.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}' + name):
                for child in list(node): node.remove(child)
                etree.SubElement(node, '{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr', val=ACCENT)
        part._blob = etree.tostring(theme)
prs.save(TARGET)
STAMP.write_text(json.dumps({'inputs': fingerprints, 'output': hashlib.sha256(TARGET.read_bytes()).hexdigest(), 'slides': len(prs.slides)}, indent=2)+'\n')
print(f'Built {len(prs.slides)} illustrated slides for {len(catalog)} projects.')
